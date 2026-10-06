from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

headers = [
    "dataset_id",
    "dataset_name",
    "source_agency",
    "exact_url",
    "access_date",
    "study_area",
    "temporal_start",
    "temporal_end",
    "temporal_resolution",
    "spatial_resolution",
    "original_format",
    "original_crs",
    "raw_file_name",
    "target_folder",
    "processing_status",
    "licence",
    "checksum",
    "data_quality",
    "responsible_person",
    "notes"
]

# নতুন Excel workbook তৈরি
workbook = Workbook()

# Active sheet নেওয়া এবং নাম পরিবর্তন
sheet = workbook.active
sheet.title = "data_inventory"

# Row 1-এ headings লেখা
for column_number, heading in enumerate(headers, start=1):
    cell = sheet.cell(
        row=1,
        column=column_number,
        value=heading
    )

    cell.font = Font(
        bold=True,
        color="FFFFFF"
    )

    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

# Header row-এর height
sheet.row_dimensions[1].height = 32

# Column width
widths = [
    14, 24, 22, 35, 14,
    20, 16, 16, 20, 20,
    18, 18, 25, 30, 20,
    18, 32, 20, 22, 35
]

for column_number, width in enumerate(widths, start=1):
    column_letter = get_column_letter(column_number)
    sheet.column_dimensions[column_letter].width = width

# প্রথম row freeze করা
sheet.freeze_panes = "A2"

# Filter যোগ করা
sheet.auto_filter.ref = "A1:T1000"

# Date column format
for row_number in range(2, 1001):
    sheet[f"E{row_number}"].number_format = "yyyy-mm-dd"
    sheet[f"G{row_number}"].number_format = "yyyy-mm-dd"
    sheet[f"H{row_number}"].number_format = "yyyy-mm-dd"

# Processing status dropdown
status_dropdown = DataValidation(
    type="list",
    formula1=(
        '"Not Started,Downloaded,Under Review,'
        'Processing,Processed,Quality Checked,Complete"'
    ),
    allow_blank=True
)

sheet.add_data_validation(status_dropdown)
status_dropdown.add("O2:O1000")

# Data quality dropdown
quality_dropdown = DataValidation(
    type="list",
    formula1='"Unknown,Poor,Fair,Good,Very Good"',
    allow_blank=True
)

sheet.add_data_validation(quality_dropdown)
quality_dropdown.add("R2:R1000")

# Excel file-এর নাম
file_name = "data_inventory_v01.xlsx"

# বর্তমান folder-এ save হবে
workbook.save(file_name)

print("Excel file সফলভাবে তৈরি হয়েছে:")
print(file_name)